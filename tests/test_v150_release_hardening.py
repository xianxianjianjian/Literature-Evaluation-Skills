from __future__ import annotations

import copy
import csv
import hashlib
import io
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import pymupdf
import artifact_binding as binding
import main_source_index
import render_exact_mirror
import run_contract_tests
import validate_a_text_layer as a
import validate_c_review as c
import validate_deliverables as d
import workflow_state
import terminology_registry as terms
import export_b_pdf as b_export
import validate_deep_reading_package as b_validate
import zotero_bridge
import zotero_attachment_contract as zotero_contract
import init_workspace
from tests.v150_test_support import bind_c_map
from tests.v150_test_support import b_report
from tests import test_v150_contracts as contracts


class ACompletenessTests(unittest.TestCase):
    def package(self, root, target, extracted, *, residual="", second=None, overlapping=False):
        pdf = root / "A.pdf"
        font = pymupdf.Font("cjk")  # portable embedded synthetic CJK font; not SimSun
        frames = [{"frame_id": "F1", "bbox_pt": [40, 650, 560, 730], "translation_action": "TRANSLATE", "role": "BODY"}]
        rows = [{"frame_ids": ["F1"], "source_text": "Memory improved REM PSG DOI p", "translated_text": target}]
        with pymupdf.open() as doc:
            page = doc.new_page(width=600, height=800)
            page.insert_font(fontname="synthetic", fontbuffer=font.buffer)
            for line_no, line in enumerate(extracted.splitlines()):
                page.insert_text((50, 90 + line_no * 15), line, fontname="synthetic", fontsize=10)
            if residual:
                page.insert_text((50, 135), residual, fontsize=10, render_mode=3)
            if second is not None:
                bbox = [40, 650, 560, 730] if overlapping else [40, 450, 560, 530]
                frames.append({"frame_id": "F2", "bbox_pt": bbox, "translation_action": "TRANSLATE", "role": "CAPTION"})
                rows.append({"frame_ids": ["F2"], "source_text": "Caption", "translated_text": target})
                if second:
                    page.insert_text((50, 90 if overlapping else 290), second, fontname="synthetic", fontsize=10)
            doc.save(pdf)
        (root / "source_inventory.json").write_text(json.dumps({"pages": [{"output_page": 1, "frame_ids": [f["frame_id"] for f in frames]}]}))
        (root / "text_frame_inventory.jsonl").write_text("\n".join(json.dumps(f) for f in frames))
        (root / "translation_ledger.jsonl").write_text("\n".join(json.dumps(r) for r in rows))
        return pdf

    def assert_gate(self, checks, code, passed, frame=None):
        check = next(check for check in checks if check.code == code)
        self.assertEqual(check.passed, passed, check)
        if frame:
            self.assertIn(frame, check.detail)
            self.assertIn(frame, check.failed_frame_ids)

    def test_all_characters_digits_statistics_and_repetitions_required(self):
        target = "睡眠后记忆得到改善并且巩固了情绪记忆 REM p = .032 [12] ± 5"
        bad = [target[:4], target[:6], target[:-8], target.replace("巩固", ""), target.replace("绪", ""),
               target.replace(".032", ".03"), target.replace("12", "1"), target.replace("REM", ""),
               target.replace("±", ""), target.replace("巩固了情绪记忆", "情绪记忆巩固了"), target + "记忆"]
        for text in bad:
            with self.subTest(text=text), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                checks = a.validate(root, self.package(root, target, text))
                self.assert_gate(checks, "A_COPY_TEXT_FIDELITY_PASS", False, "F1")
                self.assert_gate(checks, "A_READING_ORDER_PASS", False, "F1")

    def test_full_mixed_text_and_whitelisted_layout_equivalences(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            target = "睡眠记忆 REM PSG DOI F(1, 38) p = .032 [12] ± 5"
            pdf = self.package(root, target, target.replace("记忆", "记忆\n"))
            checks = a.validate(root, pdf)
            self.assertTrue(all(check.passed for check in checks), checks)
        self.assertEqual(a.normalize_text("ﬂ Ｆ（１，３８） 10⁻³"), "flF(1,38)10-3")
        self.assertEqual(a.normalize_text("记\u00ad忆 p = .032 − 5"), a.normalize_text("记忆p=.032-5"))
        self.assertNotEqual(a.normalize_text("p=.032"), a.normalize_text("p=.03"))

    def test_single_hidden_word_and_fragmented_word_are_residuals(self):
        for residual in ("Memory", "M e m o r y", "Memory improved"):
            with self.subTest(residual=residual), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                checks = a.validate(root, self.package(root, "睡眠记忆 REM PSG DOI", "睡眠记忆 REM PSG DOI", residual=residual))
                self.assert_gate(checks, "A_NO_STALE_ENGLISH_TEXT_LAYER_PASS", False, "F1")
        self.assertEqual(a._residual_tokens("PSG REM DOI", "PSG REM DOI 中文", "PSG REM DOI", {}, {"role": "BODY"}), [])
        self.assertEqual(a._residual_tokens("Memory", "中文", "Memory", {}, {"role": "FIGURE_INTERNAL"}), [])

    def test_repeated_frames_cannot_borrow_glyphs(self):
        for overlap in (False, True):
            with self.subTest(overlap=overlap), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                checks = a.validate(root, self.package(root, "睡眠记忆得到改善", "睡眠记忆得到改善", second="", overlapping=overlap))
                self.assert_gate(checks, "A_COPY_TEXT_FIDELITY_PASS", False, "F2")

    def test_missing_tounicode_fails_with_frame_page_evidence(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pdf = self.package(root, "睡眠记忆", "睡眠记忆")
            with pymupdf.open(pdf) as doc:
                xref = doc[0].get_fonts(full=True)[0][0]
                doc.xref_set_key(xref, "ToUnicode", "null")
                broken = root / "broken.pdf"
                doc.save(broken)
            checks = a.validate(root, broken)
            self.assert_gate(checks, "A_FONT_UNICODE_MAPPING_PASS", False, "F1")
            self.assertIn("page-1", next(check.detail for check in checks if check.code == "A_FONT_UNICODE_MAPPING_PASS"))

    def test_main_si_multicolumn_headings_captions_formula_and_protected_table(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            frames, rows, pages = [], [], []
            pdf = root / "mixed.pdf"
            font = pymupdf.Font("cjk")
            with pymupdf.open() as doc:
                for page_no, source_id in ((1, "SRC-M1"), (2, "SRC-S1")):
                    page = doc.new_page(width=600, height=800)
                    page.insert_font(fontname="synthetic", fontbuffer=font.buffer)
                    ids = []
                    for fid, bbox, point, target, role in (
                        (f"H{page_no}", [30, 740, 570, 775], (40, 45), "标题：主文" if page_no == 1 else "标题：补充材料", "H1"),
                        (f"L{page_no}", [30, 650, 290, 715], (40, 100), "实验 REM p=.032 F(1,38)=5.2", "BODY"),
                        (f"R{page_no}", [310, 650, 570, 715], (320, 100), "结果 PSG 95% CI [12] ± 5", "BODY"),
                        (f"C{page_no}", [30, 500, 290, 550], (40, 270), "图注：变化量及统计范围", "CAPTION"),
                        (f"T{page_no}", [310, 500, 570, 550], (320, 270), "表格：均值 38 μV", "TABLE_CELL"),
                    ):
                        page.insert_text(point, target, fontname="synthetic", fontsize=8)
                        frames.append({"frame_id": fid, "bbox_pt": bbox, "translation_action": "TRANSLATE", "role": role})
                        rows.append({"frame_ids": [fid], "source_text": "Memory improved REM PSG CI", "translated_text": target})
                        ids.append(fid)
                    page.draw_rect(pymupdf.Rect(40, 160, 200, 230), fill=(0.6, 0.8, 1))
                    page.insert_text((45, 200), "Memory figure label", fontsize=9)
                    page.insert_text((40, 700), "Journal reference DOI 10.1234/synthetic", fontsize=9)
                    pages.append({"source_id": source_id, "output_page": page_no, "frame_ids": ids})
                doc.save(pdf)
            (root / "source_inventory.json").write_text(json.dumps({"pages": pages}))
            (root / "text_frame_inventory.jsonl").write_text("\n".join(json.dumps(f) for f in frames))
            (root / "translation_ledger.jsonl").write_text("\n".join(json.dumps(r) for r in rows))
            checks = a.validate(root, pdf)
            self.assertTrue(all(check.passed for check in checks), checks)
            self.assertEqual(len(next(check.frame_statistics for check in checks if check.code == "A_COPY_TEXT_FIDELITY_PASS")), 10)

    def test_geometry_completeness_and_copy_order_are_separate(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            target = "睡眠记忆\n得到改善"
            pdf = self.package(root, target, target)
            with pymupdf.open(pdf) as original, pymupdf.open() as reordered:
                page = reordered.new_page(width=600, height=800)
                page.insert_font(fontname="synthetic", fontbuffer=pymupdf.Font("cjk").buffer)
                page.insert_text((50, 105), "得到改善", fontname="synthetic", fontsize=10)
                page.insert_text((50, 90), "睡眠记忆", fontname="synthetic", fontsize=10)
                reordered.save(root / "order.pdf")
            checks = a.validate(root, root / "order.pdf")
            self.assert_gate(checks, "A_COPY_TEXT_FIDELITY_PASS", True)
            self.assert_gate(checks, "A_READING_ORDER_PASS", False, "F1")

    def test_separate_superscript_pdf_run_keeps_formula_in_sequence(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            target = "统计量为 10⁻³"
            self.package(root, target, target)
            font = pymupdf.Font("cjk")
            with pymupdf.open() as pdf:
                page = pdf.new_page(width=600, height=800)
                page.insert_font(fontname="synthetic", fontbuffer=font.buffer)
                base = "统计量为 10"
                page.insert_text((50, 100), base, fontname="synthetic", fontsize=10)
                x = 50 + font.text_length(base, fontsize=10)
                page.insert_text((x, 96), "-3", fontname="synthetic", fontsize=6.5)
                pdf.save(root / "scientific.pdf")
            checks = a.validate(root, root / "scientific.pdf")
            self.assertTrue(all(check.passed for check in checks), checks)

    def test_redaction_preserves_art_labels_crop_and_rotation(self):
        for rotation, cropped in ((0, False), (90, False), (0, True)):
            with self.subTest(rotation=rotation, cropped=cropped), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                source = root / "source.pdf"
                with pymupdf.open() as doc:
                    page = doc.new_page(width=600, height=800)
                    page.insert_text((60, 100), "Memory improves")
                    page.insert_text((320, 100), "PSG figure label")
                    page.draw_rect(pymupdf.Rect(310, 150, 400, 250), fill=(0, 0, 1))
                    page.draw_line((310, 150), (400, 250), color=(1, 0, 0))
                    if cropped:
                        page.set_cropbox(pymupdf.Rect(20, 20, 580, 780))
                    page.set_rotation(rotation)
                    doc.save(source)
                with pymupdf.open(source) as doc:
                    page = doc[0]
                    box = page.search_for("Memory improves")[0] + (-2, -2, 2, 2)
                    frame_box = list(box * ~page.transformation_matrix)
                    before = page.get_pixmap().samples
                    protected = page.get_pixmap(clip=pymupdf.Rect(290, 140, 430, 270) * page.rotation_matrix).samples
                frames = {"F": {"translation_action": "TRANSLATE", "bbox_pt": frame_box}}
                readers = render_exact_mirror._redacted_readers(root, {"MAIN": {"pdf_path": str(source)}},
                            [{"source_id": "MAIN", "source_page": 1, "frame_ids": ["F"]}], frames)
                from pypdf import PdfWriter
                writer = PdfWriter()
                writer.add_page(readers["MAIN"].pages[0])
                dest = root / "redacted.pdf"
                writer.write(dest)
                with pymupdf.open(dest) as doc:
                    self.assertNotIn("Memory", doc[0].get_text())
                    self.assertIn("PSG figure label", doc[0].get_text())
                    self.assertEqual(protected, doc[0].get_pixmap(clip=pymupdf.Rect(290, 140, 430, 270) * doc[0].rotation_matrix).samples)
                    self.assertNotEqual(before, doc[0].get_pixmap().samples)


class CSourceAndSignoffTests(unittest.TestCase):
    def fixture(self, root):
        markdown, source_map = contracts.CSubmissionTests().fixture()
        mapping, index = bind_c_map(root, markdown, source_map)
        return markdown, mapping, index

    def gate(self, checks, name):
        return next(x for x in checks if x.code == name)

    def test_verified_main_and_complete_synthetic_signoff_pass(self):
        with tempfile.TemporaryDirectory() as tmp:
            markdown, mapping, index = self.fixture(Path(tmp))
            self.assertTrue(all(x.passed for x in c.validate(markdown, mapping, source_index=index)))

    def test_invented_page_block_quote_section_si_and_wrong_identity_fail(self):
        cases = [("page", 999), ("block", 999), ("quote", "SI-only result not in Main"),
                 ("quote", "p = .03"), ("quote", "Memory improve"),
                 ("section", "invented section"), ("source_id", "SRC-S1"), ("sha256", "0" * 64)]
        with tempfile.TemporaryDirectory() as tmp:
            markdown, mapping, index = self.fixture(Path(tmp))
            for field, value in cases:
                with self.subTest(field=field):
                    bad = copy.deepcopy(mapping)
                    bad["paragraphs"][0]["claims"][0]["source_anchor"][field] = value
                    gate = self.gate(c.validate(markdown, bad, source_index=index), "C_SOURCE_MAP_PASS")
                    self.assertFalse(gate.passed, gate)
                    self.assertIn("TEST-1-1", gate.detail)
            for field in ("doi", "version", "paper_id", "sha256"):
                bad = copy.deepcopy(mapping)
                bad[field] = "WRONG"
                self.assertFalse(self.gate(c.validate(markdown, bad, source_index=index), "C_SOURCE_MAP_PASS").passed)

    def test_bare_pass_missing_signer_stale_comment_and_rejected_direction_require_review(self):
        with tempfile.TemporaryDirectory() as tmp:
            markdown, mapping, index = self.fixture(Path(tmp))
            cases = []
            bare = copy.deepcopy(mapping)
            bare["signoff"] = {"status": "PASS", "evidence": "reviewed"}
            cases.append((markdown, bare))
            rejected = copy.deepcopy(mapping)
            first_claim = rejected["paragraphs"][0]["claims"][0]
            reversed_text = "睡眠之后记忆表现变差，研究结果显示记忆下降。"
            reversed_markdown = markdown.replace(first_claim["text"], reversed_text, 1)
            first_claim.update({"text": reversed_text, "kind": "result"})
            rejected["signoff"]["comment_sha256"] = hashlib.sha256(c.extract_comment(reversed_markdown).encode()).hexdigest()
            rejected["signoff"]["claims"][0].update({"claim_text_sha256": hashlib.sha256(reversed_text.encode()).hexdigest(),
                "status": "FAIL", "reason": "Main says Memory improved after sleep; the Chinese assertion reverses effect direction"})
            located = c.validate(reversed_markdown, rejected, source_index=index)
            self.assertTrue(self.gate(located, "C_SOURCE_MAP_PASS").passed)
            self.assertIn("improved", first_claim["source_anchor"]["quote"])
            self.assertIn("reverses effect direction", rejected["signoff"]["claims"][0]["reason"])
            cases.append((reversed_markdown, rejected))
            missing = copy.deepcopy(mapping)
            missing["signoff"]["reviewer"] = ""
            cases.append((markdown, missing))
            for text, bad in cases:
                self.assertFalse(self.gate(c.validate(text, bad, source_index=index), "C_HUMAN_SIGNOFF_PASS").passed)
                self.assertFalse(self.gate(c.validate(text, bad, source_index=index), "CLAIM_STRENGTH_PASS").passed)
            self.assertFalse(self.gate(c.validate(markdown.replace("谨慎解释", "强因果解释"), mapping, source_index=index), "C_HUMAN_SIGNOFF_PASS").passed)

    def test_source_pdf_changed_wrong_doi_and_scanned_main_require_review(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self.fixture(root)
            path = root / "source_manifest.json"
            source = json.loads(path.read_text())
            source["doi"] = "10.1234/wrong"
            path.write_text(json.dumps(source))
            with self.assertRaisesRegex(ValueError, "identity"):
                main_source_index.build_index(path)
            source["doi"] = "10.1234/synthetic"
            (root / "Main.pdf").write_bytes((root / "Main.pdf").read_bytes() + b"\n%changed")
            path.write_text(json.dumps(source))
            with self.assertRaisesRegex(ValueError, "bytes differ"):
                main_source_index.build_index(path)
            with pymupdf.open() as doc:
                doc.new_page()
                doc.save(root / "scan.pdf")
            source["sources"][0].update({"local_path": "scan.pdf", "sha256": binding.fingerprint(root / "scan.pdf")["sha256"], "bytes": (root / "scan.pdf").stat().st_size})
            path.write_text(json.dumps(source))
            with self.assertRaisesRegex(ValueError, "REVIEW_REQUIRED"):
                main_source_index.build_index(path)

    def test_reviewed_critique_exception_is_attributed_and_signed(self):
        with tempfile.TemporaryDirectory() as tmp:
            markdown, mapping, index = self.fixture(Path(tmp))
            mapping["critique_mode"] = "author_plus_major"
            mapping["paragraphs"][0]["claims"][0].update({"kind": "evaluator_critique", "attribution": "reviewer"})
            self.assertFalse(self.gate(c.validate(markdown, mapping, source_index=index), "C_SOURCE_MAP_PASS").passed)
            approval = {"reviewer": "TEST authorized reviewer", "reviewed_at": "2026-10-14T12:00:00+08:00", "reason": "TEST user requested major critique; separate evaluator attribution"}
            mapping["critique_approval"] = approval
            mapping["signoff"]["critique_approval"] = copy.deepcopy(approval)
            self.assertTrue(all(x.passed for x in c.validate(markdown, mapping, source_index=index)))


class MandatoryRunnerTests(unittest.TestCase):
    def test_skip_is_failure_for_mandatory_suite(self):
        class DeliberateSkip(unittest.TestCase):
            def runTest(self):
                self.skipTest("synthetic missing SimSun")
        with patch("sys.stderr", io.StringIO()), patch("sys.stdout", io.StringIO()):
            passed, summary = run_contract_tests.run_suite(unittest.TestSuite([DeliberateSkip()]))
        self.assertFalse(passed)
        self.assertEqual(summary, {"expected": 1, "run": 1, "passed": 0, "failures": 0, "errors": 0, "skipped": 1})


class BContentBindingTests(unittest.TestCase):
    def fixture(self, root):
        from docx import Document
        b_path = root / "B.docx"
        doc = Document()
        with pymupdf.open() as pdf:
            page = pdf.new_page(width=600, height=800)
            font = pymupdf.Font("cjk")
            page.insert_font(fontname="synthetic", fontbuffer=font.buffer)
            y = 40
            for heading in b_export.TOP_LEVEL_HEADINGS:
                doc.add_heading(heading, level=1)
                doc.add_paragraph("正文统计证据 p = .032 REM。")
                page.insert_text((30, y), heading, fontname="synthetic", fontsize=10)
                page.insert_text((30, y + 18), "正文统计证据 p = .032 REM。", fontname="synthetic", fontsize=10)
                y += 46
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "测量表格"
            table.cell(0, 1).text = "结果数值 38"
            page.insert_text((30, y), "测量表格 结果数值 38", fontname="synthetic", fontsize=10)
            rendered = root / "rendered.pdf"
            pdf.save(rendered)
        doc.save(b_path)
        report = b_report(root, b_path, rendered)
        report_path = root / "academic.json"
        report_path.write_text(json.dumps(report))
        return b_path, rendered, report_path, report

    def test_v150_content_chain_passes_and_docx_change_invalidates_report(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            b_path, rendered, report_path, report = self.fixture(root)
            result = b_export.export(b_path, report_path, rendered, root / "archive.pdf")
            self.assertTrue(result["B_PDF_EXPORT_PASS"])
            self.assertEqual(result["pdf_binding"], binding.fingerprint(rendered))
            doc = Document(b_path)
            doc.paragraphs[1].text = "正文统计证据 p = .033 REM。"
            doc.save(b_path)
            with self.assertRaisesRegex(ValueError, "stale"):
                b_export.export(b_path, report_path, rendered, root / "invalid.pdf")

    def test_same_pdf_path_changed_bytes_and_corrupt_copy_fail(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            b_path, rendered, report_path, report = self.fixture(root)
            original = rendered.read_bytes()
            rendered.write_bytes(original + b"\n% changed render")
            with self.assertRaisesRegex(ValueError, "stale"):
                b_export.export(b_path, report_path, rendered, root / "invalid.pdf")
            rendered.write_bytes(original)
            def corrupt(source, target):
                Path(target).write_bytes(Path(source).read_bytes() + b"\n%corrupt copy")
            with patch.object(b_export.shutil, "copyfile", side_effect=corrupt):
                with self.assertRaisesRegex(ValueError, "copy differs"):
                    b_export.export(b_path, report_path, rendered, root / "invalid.pdf")

    def test_other_same_named_docx_and_stale_visual_signoff_fail(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            b_path, rendered, report_path, report = self.fixture(root)
            other = root / "other" / "B.docx"
            other.parent.mkdir()
            other.write_bytes(b_path.read_bytes())
            report["artifact"] = str(other)
            report_path.write_text(json.dumps(report))
            with self.assertRaises(ValueError):
                b_export.export(b_path, report_path, rendered, root / "invalid.pdf")
            report["artifact"] = str(b_path)
            report_path.write_text(json.dumps(report))
            qa = root / "visual_qa.json"
            qa.write_text(qa.read_text() + "\n ")
            with self.assertRaisesRegex(ValueError, "sign-off differs"):
                b_export.export(b_path, report_path, rendered, root / "invalid.pdf")

    def test_new_source_version_and_workflow_mismatch_fail(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            b_path, rendered, report_path, report = self.fixture(root)
            wrong = {"paper_id": "WRONG", "workflow_id": report["workflow_id"], "contract_version": "1.5.0"}
            with self.assertRaisesRegex(ValueError, "mismatch"):
                b_export.export(b_path, report_path, rendered, root / "invalid.pdf", identity=wrong)
            (root / "Main.pdf").write_bytes((root / "Main.pdf").read_bytes() + b"\n%updated source")
            with self.assertRaisesRegex(ValueError, "source version changed"):
                b_export.export(b_path, report_path, rendered, root / "invalid.pdf")

    def test_v142_export_without_new_hashes_remains_legacy(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            b_path, rendered, report_path, report = self.fixture(root)
            for name in ("source_docx", "rendered_pdf", "visual_qa", "source_manifest", "source_fingerprints", "contract_version"):
                report.pop(name)
            report["checks"][0]["detail"] = f"rendered by independent office engine: {rendered}"
            report_path.write_text(json.dumps(report))
            self.assertTrue(b_export.export(b_path, report_path, rendered, root / "legacy.pdf", contract_version="1.4.2")["B_PDF_EXPORT_PASS"])
            with self.assertRaises(ValueError):
                b_export.export(b_path, report_path, rendered, root / "new.pdf")

    def test_relative_reports_are_independent_of_cwd(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            b_path, rendered, report_path, report = self.fixture(root)
            for name in ("source_docx", "rendered_pdf", "visual_qa"):
                report[name]["path"] = Path(report[name]["path"]).name
            report["source_manifest"] = "source_manifest.json"
            report["artifact"] = "B.docx"
            report_path.write_text(json.dumps(report))
            other = root / "other"
            other.mkdir()
            cwd = Path.cwd()
            try:
                for directory in (root, other):
                    os.chdir(directory)
                    result = b_export.export(Path("B.docx"), Path("academic.json"), Path("rendered.pdf"), Path("archive.pdf"), workspace_root=root)
                    self.assertTrue(result["B_PDF_EXPORT_PASS"])
                    self.assertEqual(Path(result["pdf_path"]), (root / "archive.pdf").resolve())
            finally:
                os.chdir(cwd)

    def test_visual_qa_binds_all_page_images_and_body_is_not_just_titles(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            b_path, rendered, report_path, report = self.fixture(root)
            pages = root / "pages"
            pages.mkdir()
            png = pages / "page-1.png"
            Image.new("RGB", (20, 20), "white").save(png)
            qa_path = root / "visual_qa.json"
            qa = json.loads(qa_path.read_text())
            qa["pages"][0]["image"] = binding.fingerprint(png)
            qa_path.write_text(json.dumps(qa))
            checks = b_validate._validate_visual_qa(qa_path, rendered, pages, b_path=b_path, contract_version="1.5.0")
            self.assertTrue(all(x.passed for x in checks), checks)
            Image.new("RGB", (20, 20), "black").save(png)
            self.assertFalse(next(x.passed for x in b_validate._validate_visual_qa(qa_path, rendered, pages, b_path=b_path, contract_version="1.5.0") if x.code == "B:visual-qa"))
            with pymupdf.open() as pdf:
                page = pdf.new_page()
                page.insert_text((40, 40), "same filename, different contents")
                pdf.save(root / "unrelated.pdf")
            self.assertTrue(b_validate.body_consistency(b_path, root / "unrelated.pdf"))

    def archive_manifest(self, root):
        b_path, rendered, report_path, report = self.fixture(root)
        exported = b_export.export(b_path, report_path, rendered, root / "archive.pdf")
        (root / "export.json").write_text(json.dumps(exported))
        manifest = workflow_state.initial_manifest("2026-W39", "synthetic-workflow")
        manifest["paper_id"] = report["paper_id"]
        manifest["outputs"]["B"].update({"local_path": "B.docx", "archive_pdf": "archive.pdf", "validation_report": "academic.json", "zotero_attachment_key": "B"})
        manifest["outputs"]["A"]["zotero_attachment_key"] = "A"
        manifest["archive"].update({"b_pdf_export_report": "export.json", "verification_report_path": "zotero.json", "zotero_parent_key": "P", "zotero_collection_key": "COL", "zotero_main_attachment_key": "MAIN"})
        pdf_binding = binding.fingerprint(root / "archive.pdf")
        verified = {"passed": True, "parent_key": "P", "collection_key": "COL", "attachment_keys": {"pdf": "MAIN", "A翻译": "A", "B精读": "B"},
                    **binding.workflow_identity(manifest), "source_fingerprints": report["source_fingerprints"],
                    "content_bindings": {"B": {"local": pdf_binding, "actual": pdf_binding, "api_md5": "TEST-md5", "actual_md5": "TEST-md5"}}}
        (root / "zotero.json").write_text(json.dumps(verified))
        return manifest, verified

    def test_archive_manifest_closes_hashes_zotero_identity_and_rejects_replacements(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest, verified = self.archive_manifest(root)
            for name in ("archive:B-PDF-export", "archive:attachment-contract"):
                self.assertTrue(next(x.passed for x in d._archive_completion_checks(manifest, workspace_root=root) if x.name == name))
            verified["content_bindings"]["B"]["actual"]["sha256"] = "0" * 64
            (root / "zotero.json").write_text(json.dumps(verified))
            self.assertFalse(next(x.passed for x in d._archive_completion_checks(manifest, workspace_root=root) if x.name == "archive:attachment-contract"))
            (root / "archive.pdf").write_bytes((root / "archive.pdf").read_bytes() + b"\n%changed")
            self.assertFalse(next(x.passed for x in d._archive_completion_checks(manifest, workspace_root=root) if x.name == "archive:B-PDF-export"))


class LegacyAndPathTests(unittest.TestCase):
    def test_legacy_c_uses_recorded_contract_not_upgraded_profile(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "knowledge").mkdir()
            profile = {"weekly_review": {"minimum_comment_chars": 800, "maximum_comment_chars": 1000, "reviewer_name": None}}
            (root / "knowledge/submission_profile.yaml").write_text(json.dumps(profile))
            c_path = root / "C.md"
            fields = "\n".join(f"## {h}\n内容" for h in c.REQUIRED_HEADINGS[:-1])
            for count, expected in ((500, True), (799, True), (499, False)):
                c_path.write_text(fields + "\n## 评论\n" + "中" * count, encoding="utf-8")
                checks = d.check_c(root, c_path, True, None, contract_version="1.4.2")
                self.assertEqual(next(x.passed for x in checks if x.name == "comment_chars"), expected)
            c_path.write_text(fields + "\n## 评论\n" + "中" * 700, encoding="utf-8")
            historical = {"submission_contract": {"weekly_review": {"minimum_comment_chars": 750}}}
            self.assertFalse(next(x.passed for x in d.check_c(root, c_path, True, None, contract_version="1.4.2", manifest_data=historical) if x.name == "comment_chars"))

    def test_v150_length_boundaries_do_not_weaken(self):
        fields = "\n".join(f"## {h}\n内容" for h in c.REQUIRED_HEADINGS[:-1])
        for count, expected in ((799, False), (800, True), (1000, True), (1001, False)):
            chunks = ["中" * (count // 4)] * 3 + ["中" * (count - 3 * (count // 4))]
            checks = c.validate(fields + "\n## 评论\n" + "\n\n".join(chunks), None)
            self.assertEqual(next(x.passed for x in checks if x.code == "COMMENT_LENGTH_PASS"), expected)

    def test_all_manifests_read_only_and_legacy_archive_unchanged(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for version in ("1.4.1", "1.4.2", "1.5.0"):
                manifest = workflow_state.initial_manifest("2026-W39")
                manifest["contract_version"] = version
                path = root / f"manifest-{version}.yaml"
                path.write_text(json.dumps(manifest))
                original = path.read_bytes()
                self.assertEqual(workflow_state.load_manifest(path)["contract_version"], version)
                self.assertEqual(path.read_bytes(), original)
                checks = d._archive_completion_checks(manifest)
                self.assertEqual(any(x.name == "archive:B-PDF-export" for x in checks), version == "1.5.0")
            real_legacy = ROOT / "weekly_reviews/2026/2026-W36/workflow_manifest.yaml"
            original = real_legacy.read_bytes()
            parsed = workflow_state.load_manifest(real_legacy)
            self.assertEqual(parsed["contract_version"], "1.4.2")
            self.assertTrue(all(x.passed for x in d._archive_completion_checks(parsed)))
            self.assertEqual(real_legacy.read_bytes(), original)

    def test_legacy_terminology_two_write_reload_cycles_preserve_evidence(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "terms.csv"
            row = dict.fromkeys(terms.LEGACY_FIELDS, "")
            row.update({"Term_ID": "TERM-0001", "English_Term": "REM", "Preferred_Chinese": "快速眼动睡眠", "Discipline": "Sleep",
                        "Subfield": "Memory", "Confidence": "HIGH", "Status": "ACTIVE", "Evidence_Level": "TE7",
                        "Evidence_IDs": "TERMEV-0001;TERMEV-0002", "Notes": "historical evidence unchanged"})
            with path.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=terms.LEGACY_FIELDS)
                writer.writeheader()
                writer.writerow(row)
            expected = terms.read_registry(path)
            for _ in range(2):
                terms.write_registry(path, expected)
                self.assertEqual(terms.read_registry(path), expected)
            self.assertEqual(expected[0]["Evidence_IDs"], row["Evidence_IDs"])
            self.assertEqual(expected[0]["Domain"], row["Discipline"])

    def test_init_preserves_historical_profile_and_manifest(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            init_workspace.initialize_workspace(root, template_root=ROOT / "assets/workspace-template")
            profile = root / "knowledge/submission_profile.yaml"
            profile.write_text('{"weekly_review":{"minimum_comment_chars":650}}')
            manifest = root / "weekly_reviews/old.yaml"
            manifest.write_text('{"contract_version":"1.4.2"}')
            originals = {p: p.read_bytes() for p in (profile, manifest)}
            init_workspace.initialize_workspace(root, template_root=ROOT / "assets/workspace-template")
            self.assertEqual({p: p.read_bytes() for p in originals}, originals)

    def test_traversal_external_missing_and_symlink_escape_rejected(self):
        import subprocess
        with tempfile.TemporaryDirectory() as tmp, tempfile.TemporaryDirectory() as external:
            root, outside = Path(tmp), Path(external)
            target = outside / "data.json"
            target.write_text("{}")
            for value in ("../data.json", str(target), "missing.json"):
                with self.assertRaises(ValueError):
                    binding.resolve_path(value, root)
            self.assertEqual(binding.resolve_path(target, root, allowed_roots=(outside,)), target.resolve())
            link = root / "link"
            try:
                link.symlink_to(outside, target_is_directory=True)
            except OSError:
                if os.name != "nt":
                    raise
                result = subprocess.run(["cmd", "/c", "mklink", "/J", str(link), str(outside)], capture_output=True)
                self.assertEqual(result.returncode, 0, result.stderr)
            with self.assertRaisesRegex(ValueError, "outside"):
                binding.resolve_path(link / "data.json", root)


class ZoteroContentTests(unittest.TestCase):
    def fixture(self, root):
        import argparse
        from urllib.request import pathname2url
        local = root / "local.pdf"
        actual = root / "uploaded.pdf"
        with pymupdf.open() as pdf:
            pdf.new_page()
            pdf.save(local)
        actual.write_bytes(local.read_bytes())
        parent = {"key": "P", "data": {"itemType": "journalArticle", "DOI": "10.1234/test", "collections": ["COL"]}}
        children = [{"key": role, "data": {"itemType": "attachment", "parentItem": "P", "title": zotero_contract.role_title(role),
                                           "contentType": "application/pdf", "filename": "uploaded.pdf", "md5": hashlib.md5(actual.read_bytes()).hexdigest()}}
                    for role in ("MAIN", "A", "B")]
        spec = {"parent_key": "P", "doi": "10.1234/test", "collection_key": "COL", "paper_id": "TEST", "workflow_id": "TEST-WF", "contract_version": "1.5.0",
                "source_fingerprints": {"SRC-M1": binding.fingerprint(local)},
                "attachments": [{"role": role, "key": role, "mime": "application/pdf", "local_path": "local.pdf", "local": binding.fingerprint(local)} for role in ("MAIN", "A", "B")]}
        spec_path = root / "spec.json"
        spec_path.write_text(json.dumps(spec))
        args = argparse.Namespace(spec=spec_path, parent_key="P", workspace_root=root, allow_external_root=[],
                                  api_base_url="http://127.0.0.1:23119/api", library_prefix="users/0", timeout=3, report=root / "verification.json")
        url = actual.as_uri()
        return parent, children, spec, args, actual, url

    def test_live_read_bytes_bind_local_sha_and_api_md5(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            parent, children, spec, args, actual, url = self.fixture(root)
            with patch.object(zotero_bridge, "_parent_children", return_value=(parent, children)), patch.object(zotero_bridge, "request_text", return_value=url), patch.object(zotero_bridge, "emit"):
                self.assertEqual(zotero_bridge.command_verify_archive(args), 0)
            result = json.loads(args.report.read_text(encoding="utf-8"))
            self.assertEqual(result["content_bindings"]["B"]["actual"], binding.fingerprint(actual))
            self.assertEqual(result["content_bindings"]["B"]["api_md5"], hashlib.md5(actual.read_bytes()).hexdigest())
            actual.write_bytes(actual.read_bytes() + b"\n% same name different uploaded bytes")
            with patch.object(zotero_bridge, "_parent_children", return_value=(parent, children)), patch.object(zotero_bridge, "request_text", return_value=url), patch.object(zotero_bridge, "emit"):
                self.assertEqual(zotero_bridge.command_verify_archive(args), 2)
            self.assertFalse(json.loads(args.report.read_text(encoding="utf-8"))["passed"])

    def test_local_fingerprint_and_parent_collection_mismatches_fail(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            parent, children, spec, args, actual, url = self.fixture(root)
            spec["attachments"][2]["local"]["sha256"] = "0" * 64
            result = zotero_contract.validate_archive(parent, children, spec, workspace_root=root)
            self.assertFalse(result["passed"])
            self.assertTrue(any("local attachment differs" in x for x in result["failures"]))
            parent["data"]["collections"] = []
            result = zotero_contract.validate_archive(parent, children, spec, workspace_root=root)
            self.assertIn("parent collection mismatch", result["failures"])


if __name__ == "__main__":
    unittest.main()
