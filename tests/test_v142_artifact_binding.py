from __future__ import annotations

import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'scripts'))
import deep_reading_evidence as deep
import audit_untranslated_residuals as residual
import extract_text_frames_impl as extractor
import sanitize_docx_metadata as metadata
import source_package
import validate_deliverables as deliverables
import validate_deep_reading_package as bvalidator


class ArtifactBindingTests(unittest.TestCase):
    def test_sanitizer_preserves_ignorable_namespaces(self):
        xml = b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" mc:Ignorable="w14"><w:body><w:p/></w:body></w:document>'
        result = metadata._remove_comment_markup(xml)
        self.assertIn(b'xmlns:w14=', result)
        self.assertIn(b'mc:Ignorable="w14"', result)

    def test_declaration_does_not_hide_bad_ignorable_prefix(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'bad.docx'
            with zipfile.ZipFile(path,'w') as z:
                z.writestr('word/document.xml', '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" mc:Ignorable="w14"><w:body/></w:document>')
            self.assertTrue(bvalidator._relationship_checks(path))

    def test_prose_in_table_does_not_close_narrative(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            base=Path(tmp); path=base/'b.docx'
            prose='核心解释必须存在于正常段落，而不是仅在表格中。'
            doc=Document(); doc.add_table(rows=1,cols=1).cell(0,0).text=prose; doc.save(path)
            (base/'source_figure_inventory.jsonl').write_text('',encoding='utf-8')
            narrative={'sections':{'RESULTS':{'prose':prose}},'claim_closure':[]}
            self.assertFalse(deep.validate_docx_evidence(path,narrative,{'figures':[]},base)['passed'])
            doc.add_paragraph(prose); doc.save(path)
            self.assertTrue(deep.validate_docx_evidence(path,narrative,{'figures':[]},base)['passed'])

    def test_unembedded_file_does_not_count_as_figure(self):
        from docx import Document
        from PIL import Image
        with tempfile.TemporaryDirectory() as tmp:
            base=Path(tmp); path=base/'b.docx'; img=base/'f.png'
            Image.new('RGB',(10,10),'white').save(img)
            caption='图1 原图的完整中文图注。'; explanation='原图显示估计量及其不确定性，不能独立确立因果方向。'
            doc=Document(); doc.add_paragraph(caption); doc.add_paragraph(explanation); doc.add_paragraph('SRC-M1 Fig.1'); doc.save(path)
            (base/'source_figure_inventory.jsonl').write_text(json.dumps({'figure_id':'F1','translated_text':caption}),encoding='utf-8')
            figure={'figure_id':'F1','category':'CORE_RESULT','image_path':'f.png','translated_caption':caption,'interpretation':explanation,'source_anchor':'SRC-M1 Fig.1'}
            narrative={'sections':{},'claim_closure':[]}
            self.assertFalse(deep.validate_docx_evidence(path,narrative,{'figures':[figure]},base)['passed'])
            doc.add_picture(str(img)); doc.save(path)
            self.assertTrue(deep.validate_docx_evidence(path,narrative,{'figures':[figure]},base)['passed'])

    def test_omitted_source_figure_fails(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            base=Path(tmp); path=base/'b.docx'; Document().save(path)
            (base/'source_figure_inventory.jsonl').write_text(json.dumps({'figure_id':'F2'}),encoding='utf-8')
            self.assertFalse(deep.validate_docx_evidence(path,{'sections':{},'claim_closure':[]},{'figures':[]},base)['passed'])

    def test_short_source_word_is_not_ignored(self):
        self.assertEqual(residual._shared_source_runs('of','effect of sleep',set()),['of'])

    def test_url_does_not_hide_adjacent_prose(self):
        self.assertTrue(residual._shared_source_runs('memory https://example.org','memory https://example.org',set()))

    def test_new_contract_has_no_global_paper_whitelist(self):
        self.assertNotIn('memory',residual._allowed_words([],legacy=False))
        self.assertNotIn('e-wake',residual._allowed_words([],legacy=False))

    def test_compound_allowance_requires_reason(self):
        self.assertNotIn('e-wake',residual._allowed_words([{'untranslated_tokens':[{'text':'E-Wake'}]}],legacy=False))
        self.assertIn('e-wake',residual._allowed_words([{'untranslated_tokens':[{'text':'E-Wake','reason':'experimental group code'}]}],legacy=False))

    def test_low_confidence_ocr_review_requires_complete_human_record(self):
        self.assertFalse(residual._valid_ocr_review({}))
        self.assertFalse(residual._valid_ocr_review({'ocr_review': {'finding': 'NO_UNTRANSLATED_PROSE'}}))
        self.assertTrue(residual._valid_ocr_review({'ocr_review': {
            'finding': 'NO_UNTRANSLATED_PROSE', 'reviewer': 'translation QC',
            'reviewed_date': '2026-09-09', 'method': 'source/output crop side-by-side',
        }}))

    def test_pdf_subset_font_emphasis(self):
        self.assertEqual(extractor.font_emphasis('ABCDE+AdvOT34fe.B'),(True,False))
        self.assertEqual(extractor.font_emphasis('ABCDE+AdvOT34fe.I'),(False,True))

    def test_body_with_link_is_still_translatable(self):
        self.assertEqual(extractor._role('body',10,10,'Data are available at https://example.org.'),'BODY')
        self.assertEqual(extractor._role('body',10,10,'https://example.org'),'DOI_URL')

    def test_main_si_reference_discovery(self):
        found=source_package.discover_si_references('See Supplementary Table 3 and Supplementary Figure 1. Supplementary Methods provide details.')
        self.assertEqual(set(found),{'Supplementary Table 3','Supplementary Figure 1','Supplementary Methods'})

    def test_deliverable_validator_forwards_b_contract(self):
        with patch.object(deliverables.deep_reading_validator,'validate_package',return_value=[]) as check:
            deliverables.check_deep_reading_package(Path('.'),Path('b.docx'),True,None,contract_version='1.4.2')
            self.assertEqual(check.call_args.kwargs['contract_version'],'1.4.2')

    def test_deliverable_validator_forwards_a_contract(self):
        with patch.object(deliverables.translation_validator,'validate_package_detailed',return_value=([],None,'EXACT_TEXT_FRAME')) as check:
            deliverables.check_translation_package(Path('.'),Path('a.pdf'),'FULL_MIRROR',True,contract_version='1.4.2')
            self.assertEqual(check.call_args.args[-1],'1.4.2')

if __name__=='__main__': unittest.main()
