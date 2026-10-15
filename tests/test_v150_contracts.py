from __future__ import annotations

import io
import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from pypdf import PdfWriter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

import export_b_pdf
import render_exact_mirror
import validate_a_text_layer
import validate_c_review
import validate_deliverables
import workflow_state
import zotero_attachment_contract as zotero_contract


class ATextLayerTests(unittest.TestCase):
    def test_search_keyword_does_not_join_across_punctuation_or_identifiers(self) -> None:
        self.assertEqual(validate_a_text_layer._search_keyword("图 S2 刺激类型"), "刺激类型")
        self.assertEqual(validate_a_text_layer._search_keyword("同样，效价评分"), "同样")

    @unittest.skipUnless(render_exact_mirror.pymupdf, "PyMuPDF unavailable")
    def test_two_column_source_english_is_removed_before_chinese_overlay(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.pdf"
            writer = canvas.Canvas(str(source), pagesize=(600, 800))
            writer.setFont("Helvetica", 11)
            writer.drawString(50, 700, "Memory improves after sleep")
            writer.drawString(320, 700, "Night cycles support recall")
            writer.drawString(50, 750, "DOI: 10.1234/example")
            writer.save()
            frames = {
                "left": {"frame_id": "left", "bbox_pt": [40, 680, 280, 720], "translation_action": "TRANSLATE"},
                "right": {"frame_id": "right", "bbox_pt": [310, 680, 570, 720], "translation_action": "TRANSLATE"},
            }
            pages = [{"source_id": "SRC-M1", "source_page": 1, "output_page": 1,
                      "frame_ids": ["left", "right"]}]
            inventory = {"sources": [{"source_id": "SRC-M1", "pdf_path": str(source)}], "pages": pages}
            (root / "source_inventory.json").write_text(json.dumps(inventory), encoding="utf-8")
            (root / "text_frame_inventory.jsonl").write_text(
                "\n".join(json.dumps(row) for row in frames.values()), encoding="utf-8"
            )
            ledger = [
                {"frame_ids": ["left"], "source_text": "Memory improves after sleep", "translated_text": "睡眠后记忆得到改善"},
                {"frame_ids": ["right"], "source_text": "Night cycles support recall", "translated_text": "睡眠周期有助于回忆"},
            ]
            (root / "translation_ledger.jsonl").write_text(
                "\n".join(json.dumps(row, ensure_ascii=False) for row in ledger), encoding="utf-8"
            )
            font = Path(r"C:\Windows\Fonts\simsun.ttc")
            if not font.is_file():
                self.skipTest("SimSun unavailable")
            pdfmetrics.registerFont(TTFont("SimSun-v150-test", str(font), subfontIndex=0))
            overlay = io.BytesIO()
            drawing = canvas.Canvas(overlay, pagesize=(600, 800))
            drawing.setFont("SimSun-v150-test", 11)
            drawing.drawString(50, 700, "睡眠后记忆得到改善")
            drawing.drawString(320, 700, "睡眠周期有助于回忆")
            drawing.save()
            from pypdf import PdfReader
            overlay.seek(0)
            redacted = render_exact_mirror._redacted_readers(
                root, {"SRC-M1": inventory["sources"][0]}, pages, frames
            )
            output = PdfWriter()
            output.add_page(redacted["SRC-M1"].pages[0])
            output.pages[0].merge_page(PdfReader(overlay).pages[0])
            a_pdf = root / "A.pdf"
            with a_pdf.open("wb") as handle:
                output.write(handle)
            text = render_exact_mirror.pymupdf.open(a_pdf)[0].get_text()
            self.assertNotIn("Memory improves", text)
            self.assertNotIn("Night cycles", text)
            self.assertIn("10.1234/example", text)
            checks = validate_a_text_layer.validate(root, a_pdf)
            self.assertTrue(all(item.passed for item in checks), checks)
            # The original visually covered page remains extractably English.
            stale = PdfWriter()
            stale.add_page(PdfReader(str(source)).pages[0])
            stale.pages[0].merge_page(PdfReader(overlay).pages[0])
            stale_path = root / "stale.pdf"
            with stale_path.open("wb") as handle:
                stale.write(handle)
            stale_checks = validate_a_text_layer.validate(root, stale_path)
            self.assertFalse(next(item.passed for item in stale_checks if item.code == "A_NO_STALE_ENGLISH_TEXT_LAYER_PASS"))


class CSubmissionTests(unittest.TestCase):
    def test_tabarak_submission_regression_fixture(self) -> None:
        fixture = ROOT / "tests" / "fixtures" / "v150"
        markdown = (fixture / "tabarak_c.md").read_text(encoding="utf-8")
        source_map = json.loads((fixture / "tabarak_c_source_map.json").read_text(encoding="utf-8"))
        checks = validate_c_review.validate(markdown, source_map)
        # This historical fixture has no actual Main or attributable sign-off.
        # Its public format remains a regression, but it cannot prove sources.
        self.assertFalse(next(item.passed for item in checks if item.code == "C_SOURCE_MAP_PASS"))
        self.assertTrue(all(item.passed for item in checks[:4]), checks)
        count = validate_c_review.effective_visible_char_count(validate_c_review.extract_comment(markdown))
        self.assertTrue(850 <= count <= 950, count)
        comment = validate_c_review.extract_comment(markdown)
        self.assertIn("第二睡眠周期", comment)
        self.assertIn("整夜均值转向时间演变", comment)
        self.assertNotRegex(comment, r"AUD-|B\s*§|SI\s*(?:Table|Figure)")
        self.assertIn("不是直接提供临床治疗结论", comment)

    def fixture(self) -> tuple[str, dict]:
        base = "研究围绕睡眠与记忆的关系提出明确问题，并用相互衔接的观察与干预分析追踪变化。作者报告的主要结果支持谨慎解释，同时指出测量和机制之间仍有未解决的联系。"
        paragraphs = [base * 3, base * 3, base * 3, base * 3]
        body = "\n\n".join(paragraphs)
        headings = "\n".join(f"## {name}\n内容\n" for name in validate_c_review.REQUIRED_HEADINGS[:-1])
        markdown = headings + "\n## 评论\n" + body
        source_map = {
            "paragraphs": [
                {"paragraph_index": index + 1, "claims": [{"text": base[:20], "source_class": "MAIN",
                  "source_anchor": "Main Discussion p.1", "kind": "limitation",
                  "author_acknowledged": True}]}
                for index in range(4)
            ],
            "semantic_review": {name: {"status": "PASS", "evidence": "Main p.1 reviewed"}
                                for name in validate_c_review.SEMANTIC_GATES},
        }
        return markdown, source_map

    def test_count_ignores_markdown_and_public_c_rejects_internal_and_si(self) -> None:
        self.assertEqual(validate_c_review.effective_visible_char_count("**中文** [ABC](https://x)"), 5)
        markdown, source_map = self.fixture()
        count = validate_c_review.effective_visible_char_count(validate_c_review.extract_comment(markdown))
        self.assertTrue(800 <= count <= 1000)
        checks = validate_c_review.validate(markdown, source_map)
        from tests.v150_test_support import bind_c_map
        with tempfile.TemporaryDirectory() as directory:
            source_map, index = bind_c_map(Path(directory), markdown, source_map)
            checks = validate_c_review.validate(markdown, source_map, source_index=index)
            self.assertTrue(all(item.passed for item in checks), checks)
        long_review = markdown + ("研究结论应保持谨慎。" * 20)
        long_checks = validate_c_review.validate(long_review, source_map)
        self.assertFalse(next(item.passed for item in long_checks if item.code == "COMMENT_LENGTH_PASS"))
        bad = markdown.replace("## 评论\n", "## 评论\nSI Table S7 与 AUD-007。")
        bad_checks = validate_c_review.validate(bad, source_map, min_chars=count, max_chars=count)
        self.assertFalse(next(item.passed for item in bad_checks if item.code == "NO_INTERNAL_AUDIT_METADATA_PASS"))
        source_map["paragraphs"][0]["claims"][0]["source_class"] = "SI"
        mapped = validate_c_review.validate(markdown, source_map, min_chars=count, max_chars=count, source_index=index)
        self.assertFalse(next(item.passed for item in mapped if item.code == "C_SOURCE_MAP_PASS"))


class ArchiveTests(unittest.TestCase):
    def test_b_pdf_promotes_only_validated_independent_render(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            b_path = root / "B.docx"
            document = Document()
            for heading in export_b_pdf.TOP_LEVEL_HEADINGS:
                document.add_heading(heading, level=1)
                document.add_paragraph("已核验的研究内容。")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "方法"
            table.cell(1, 1).text = "结果"
            document.add_page_break()
            document.save(b_path)
            font = Path(r"C:\Windows\Fonts\simsun.ttc")
            if not font.is_file():
                self.skipTest("SimSun unavailable")
            pdfmetrics.registerFont(TTFont("SimSun-v150-b", str(font), subfontIndex=0))
            rendered = root / "rendered.pdf"
            drawing = canvas.Canvas(str(rendered), pagesize=(595, 842))
            drawing.setFont("SimSun-v150-b", 12)
            y = 800
            for index, heading in enumerate(export_b_pdf.TOP_LEVEL_HEADINGS):
                if index == 6:
                    drawing.showPage()
                    drawing.setFont("SimSun-v150-b", 12)
                    y = 800
                drawing.drawString(40, y, heading)
                y -= 28
                drawing.drawString(40, y, "已核验的研究内容。")
                y -= 22
            drawing.drawString(40, y, "方法 结果")
            drawing.save()
            report_path = root / "academic.json"
            from tests.v150_test_support import b_report
            report = b_report(root, b_path, rendered)
            report_path.write_text(json.dumps(report), encoding="utf-8")
            output = root / "B-archive.pdf"
            result = export_b_pdf.export(b_path, report_path, rendered, output)
            self.assertTrue(result["B_PDF_EXPORT_PASS"])
            self.assertEqual(result["pages"], 2)
            report["passed"] = False
            report_path.write_text(json.dumps(report), encoding="utf-8")
            with self.assertRaises(ValueError):
                export_b_pdf.export(b_path, report_path, rendered, root / "bad.pdf")

    def test_role_contract_and_archive_pdf_gate(self) -> None:
        parent = {"key": "PARENT", "data": {"itemType": "journalArticle", "DOI": "10.1/test"}}
        rows = [
            {"role": role, "key": role, "mime": "application/pdf"}
            for role in ("MAIN", "A", "B")
        ]
        children = [
            {"key": role, "data": {"itemType": "attachment", "parentItem": "PARENT",
              "title": zotero_contract.role_title(role), "contentType": "application/pdf",
              "filename": f"paper-{role}.pdf", "md5": "abc"}}
            for role in ("MAIN", "A", "B")
        ]
        spec = {"parent_key": "PARENT", "doi": "10.1/test", "attachments": rows}
        self.assertTrue(zotero_contract.validate_archive(parent, children, spec)["passed"])
        self.assertFalse(zotero_contract.validate_archive(parent, children + [children[0]], spec)["passed"])
        with tempfile.TemporaryDirectory() as directory:
            manifest = workflow_state.initial_manifest("2026-W37")
            archive = manifest["archive"]
            archive.update({"zotero_parent_key": "PARENT", "zotero_collection_key": "COLLECTION",
                            "zotero_main_attachment_key": "MAIN"})
            manifest["outputs"]["A"]["zotero_attachment_key"] = "A"
            manifest["outputs"]["B"]["zotero_attachment_key"] = "B"
            archive["source_archive"].update({"status": "COMPLETE", "source_ids": ["SRC-M1"]})
            archive["output_archive"].update({"status": "COMPLETE", "output_attachment_keys": {"A": "A", "B": "B"}})
            report = Path(directory) / "zotero.json"
            report.write_text(json.dumps(zotero_contract.validate_archive(parent, children, spec)), encoding="utf-8")
            archive["verification_report_path"] = str(report)
            checks = validate_deliverables._archive_completion_checks(manifest)
            self.assertFalse(next(item.passed for item in checks if item.name == "archive:B-PDF-export"))


if __name__ == "__main__":
    unittest.main()
